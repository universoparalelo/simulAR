"""
Módulo para exportar datos de simulaciones a Excel.
"""

from datetime import datetime
from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class ExcelExporter:
    """Exportador de datos a archivos Excel."""

    @staticmethod
    def export_simulation(sim_info: Dict, output_path: str = None) -> str:
        """
        Exporta la información de una simulación a un archivo Excel.

        Args:
            sim_info: Diccionario con información de la simulación
            output_path: Ruta de salida (si no se proporciona, se usa el directorio actual)

        Returns:
            Ruta del archivo creado
        """
        if output_path is None:
            output_path = Path.cwd()
        else:
            output_path = Path(output_path)

        # Crear libro de trabajo
        wb = Workbook()

        # Hoja 1: Información General
        ws_info = wb.active
        ws_info.title = "Información"

        ExcelExporter._add_info_sheet(ws_info, sim_info)

        # Hoja 2: Archivos
        ws_files = wb.create_sheet("Archivos")
        ExcelExporter._add_files_sheet(ws_files, sim_info.get("files", []))

        # Hoja 3: Métricas (si existen)
        if "metrics" in sim_info and sim_info["metrics"]:
            ws_metrics = wb.create_sheet("Métricas")
            ExcelExporter._add_metrics_sheet(ws_metrics, sim_info["metrics"])

        # Guardar archivo
        filename = f"{sim_info.get('title', 'simulation')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = output_path / filename

        wb.save(str(filepath))
        return str(filepath)

    @staticmethod
    def _add_info_sheet(worksheet, sim_info: Dict):
        """Agrega la hoja de información general."""
        # Estilos
        header_fill = PatternFill(
            start_color="4472C4", end_color="4472C4", fill_type="solid"
        )
        header_font = Font(bold=True, color="FFFFFF")
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # Título
        worksheet.merge_cells("A1:B1")
        title_cell = worksheet["A1"]
        title_cell.value = "INFORMACIÓN DE LA SIMULACIÓN"
        title_cell.font = Font(bold=True, size=14)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        # Datos
        data = [
            ("Título", sim_info.get("title", "N/A")),
            ("Programa", sim_info.get("program", "N/A")),
            ("Ruta", sim_info.get("path", "N/A")),
            ("Fecha", sim_info.get("date", "N/A")),
            ("Tamaño (bytes)", sim_info.get("size_bytes", 0)),
            ("Descripción", sim_info.get("description", "")),
            ("Fecha de Creación", sim_info.get("created_at", "N/A")),
            ("Última Actualización", sim_info.get("updated_at", "N/A")),
        ]

        row = 3
        for label, value in data:
            # Label
            label_cell = worksheet.cell(row=row, column=1)
            label_cell.value = label
            label_cell.font = Font(bold=True)
            label_cell.fill = PatternFill(
                start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
            )
            label_cell.border = border

            # Value
            value_cell = worksheet.cell(row=row, column=2)
            value_cell.value = value
            value_cell.border = border
            value_cell.alignment = Alignment(wrap_text=True)

            row += 1

        # Ajustar ancho de columnas
        worksheet.column_dimensions["A"].width = 25
        worksheet.column_dimensions["B"].width = 50

    @staticmethod
    def _add_files_sheet(worksheet, files: List[Dict]):
        """Agrega la hoja de archivos."""
        # Estilos
        header_fill = PatternFill(
            start_color="4472C4", end_color="4472C4", fill_type="solid"
        )
        header_font = Font(bold=True, color="FFFFFF")
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # Encabezados
        headers = ["Nombre", "Tipo", "Tamaño (MB)", "Ruta"]
        for col, header in enumerate(headers, 1):
            cell = worksheet.cell(row=1, column=col)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Datos
        for row, file_info in enumerate(files, 2):
            cells_data = [
                file_info.get("filename", "N/A"),
                file_info.get("file_type", "N/A"),
                file_info.get("size_bytes", 0) / (1024 * 1024),
                file_info.get("file_path", "N/A"),
            ]

            for col, value in enumerate(cells_data, 1):
                cell = worksheet.cell(row=row, column=col)
                cell.value = value
                cell.border = border

                if col == 3:  # Columna de tamaño
                    cell.value = f"{value:.2f}"
                    cell.alignment = Alignment(horizontal="right")

        # Ajustar ancho de columnas
        worksheet.column_dimensions["A"].width = 30
        worksheet.column_dimensions["B"].width = 15
        worksheet.column_dimensions["C"].width = 15
        worksheet.column_dimensions["D"].width = 50

    @staticmethod
    def _add_metrics_sheet(worksheet, metrics: List[Dict]):
        """Agrega la hoja de métricas."""
        # Estilos
        header_fill = PatternFill(
            start_color="4472C4", end_color="4472C4", fill_type="solid"
        )
        header_font = Font(bold=True, color="FFFFFF")
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # Encabezados
        headers = ["Métrica", "Valor", "Unidad", "Calculada en"]
        for col, header in enumerate(headers, 1):
            cell = worksheet.cell(row=1, column=col)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Datos
        for row, metric in enumerate(metrics, 2):
            cells_data = [
                metric.get("metric_name", "N/A"),
                metric.get("metric_value", 0),
                metric.get("metric_unit", ""),
                metric.get("calculated_at", "N/A"),
            ]

            for col, value in enumerate(cells_data, 1):
                cell = worksheet.cell(row=row, column=col)
                cell.value = value
                cell.border = border

                if col == 2:  # Columna de valor
                    cell.alignment = Alignment(horizontal="right")

        # Ajustar ancho de columnas
        worksheet.column_dimensions["A"].width = 20
        worksheet.column_dimensions["B"].width = 15
        worksheet.column_dimensions["C"].width = 15
        worksheet.column_dimensions["D"].width = 25

    @staticmethod
    def export_multiple_simulations(
        simulations: List[Dict], output_path: str = None
    ) -> str:
        """
        Exporta múltiples simulaciones a un único archivo Excel.

        Args:
            simulations: Lista de diccionarios con información de simulaciones
            output_path: Ruta de salida

        Returns:
            Ruta del archivo creado
        """
        if output_path is None:
            output_path = Path.cwd()
        else:
            output_path = Path(output_path)

        wb = Workbook()
        wb.remove(wb.active)

        # Crear hoja resumen
        ws_summary = wb.create_sheet("Resumen")
        ExcelExporter._add_summary_sheet(ws_summary, simulations)

        # Crear hoja para cada simulación
        for sim in simulations:
            ws = wb.create_sheet(sim.get("title", "Simulación")[:31])
            ExcelExporter._add_info_sheet(ws, sim)

        # Guardar archivo
        filename = (
            f"simulaciones_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
        filepath = output_path / filename

        wb.save(str(filepath))
        return str(filepath)

    @staticmethod
    def _add_summary_sheet(worksheet, simulations: List[Dict]):
        """Agrega la hoja de resumen."""
        # Estilos
        header_fill = PatternFill(
            start_color="4472C4", end_color="4472C4", fill_type="solid"
        )
        header_font = Font(bold=True, color="FFFFFF")
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # Encabezados
        headers = ["Título", "Programa", "Fecha", "Tamaño (MB)", "Archivos"]
        for col, header in enumerate(headers, 1):
            cell = worksheet.cell(row=1, column=col)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Datos
        for row, sim in enumerate(simulations, 2):
            cells_data = [
                sim.get("title", "N/A"),
                sim.get("program", "N/A"),
                sim.get("date", "N/A"),
                sim.get("size_bytes", 0) / (1024 * 1024),
                len(sim.get("files", [])),
            ]

            for col, value in enumerate(cells_data, 1):
                cell = worksheet.cell(row=row, column=col)
                cell.value = value
                cell.border = border

                if col == 4:  # Columna de tamaño
                    cell.value = f"{value:.2f}"
                    cell.alignment = Alignment(horizontal="right")

        # Ajustar ancho de columnas
        worksheet.column_dimensions["A"].width = 25
        worksheet.column_dimensions["B"].width = 15
        worksheet.column_dimensions["C"].width = 15
        worksheet.column_dimensions["D"].width = 15
        worksheet.column_dimensions["E"].width = 12
